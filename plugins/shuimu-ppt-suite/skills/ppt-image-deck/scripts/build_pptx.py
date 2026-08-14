#!/usr/bin/env python3
"""Assemble images/NN.png into a full-bleed 16:9 PPTX — one image per slide.

Each image becomes a slide with zero margin (image fills the whole slide).
Images are ordered by the leading number in their filename (01.png, 02.png, ...).

Usage:
    python build_pptx.py --images <dir> --out <file.pptx> [--pdf] [--size 13.333x7.5]

Missing python-pptx is auto-installed (uv pip / pip). --pdf needs LibreOffice (soffice).
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
from pathlib import Path

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}


def ensure_pptx() -> None:
    try:
        import pptx  # noqa: F401
        return
    except ImportError:
        pass
    for cmd in (
        ["uv", "pip", "install", "python-pptx"],
        [sys.executable, "-m", "pip", "install", "python-pptx"],
    ):
        try:
            subprocess.run(cmd, check=True)
            import pptx  # noqa: F401
            return
        except Exception:
            continue
    sys.exit("无法自动安装 python-pptx，请手动执行： pip install python-pptx")


def natural_key(p: Path):
    """Sort by the first integer in the filename; fall back to the name."""
    m = re.search(r"\d+", p.stem)
    return (int(m.group()) if m else 1 << 30, p.name.lower())


def collect_images(images_dir: Path) -> list[Path]:
    imgs = [p for p in images_dir.iterdir()
            if p.is_file() and p.suffix.lower() in IMAGE_EXTS]
    if not imgs:
        sys.exit(f"未在 {images_dir} 找到图片（{sorted(IMAGE_EXTS)}）")
    return sorted(imgs, key=natural_key)


def build(images: list[Path], out: Path, w_in: float, h_in: float) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(w_in)
    prs.slide_height = Inches(h_in)
    blank = prs.slide_layouts[6]  # fully blank layout

    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(img), Inches(0), Inches(0),
            width=prs.slide_width, height=prs.slide_height,
        )

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"✅ 已生成 {out}  （{len(images)} 页）")
    for i, img in enumerate(images, 1):
        print(f"   {i:>3}  {img.name}")


def export_pdf(pptx_path: Path) -> None:
    import shutil
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("ℹ️  未找到 LibreOffice(soffice)，跳过 PDF 导出。")
        return
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf",
         "--outdir", str(pptx_path.parent), str(pptx_path)],
        check=True,
    )
    print(f"✅ 已导出 PDF： {pptx_path.with_suffix('.pdf')}")


def parse_size(s: str) -> tuple[float, float]:
    m = re.fullmatch(r"\s*([\d.]+)\s*[xX]\s*([\d.]+)\s*", s)
    if not m:
        sys.exit("--size 格式应为 宽x高（英寸），例如 13.333x7.5")
    return float(m.group(1)), float(m.group(2))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--images", required=True, type=Path, help="图片目录（含 01.png…）")
    ap.add_argument("--out", required=True, type=Path, help="输出 .pptx 路径")
    ap.add_argument("--size", default="13.333x7.5",
                    help="幻灯片尺寸（英寸），默认 13.333x7.5 = 16:9 宽屏")
    ap.add_argument("--pdf", action="store_true", help="同时导出 PDF（需 LibreOffice）")
    args = ap.parse_args()

    ensure_pptx()
    w_in, h_in = parse_size(args.size)
    images = collect_images(args.images)
    build(images, args.out, w_in, h_in)
    if args.pdf:
        export_pdf(args.out)


if __name__ == "__main__":
    main()
